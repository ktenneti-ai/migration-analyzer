"""Renders an AssessmentReport to downloadable PDF / DOCX bytes.

Section content is plain text built by app.reporting.blueprint, so both
renderers just walk the same ReportSection list — nothing report-specific
is decided here.
"""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.metadata.models import AssessmentReport


def render_pdf(report: AssessmentReport) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"Migration Assessment Report — {report.project_name}", styles["Title"]),
        Paragraph(f"Generated {report.generated_at.isoformat()}", styles["Normal"]),
        Spacer(1, 16),
        Paragraph("Executive Summary", styles["Heading1"]),
    ]
    for key, value in report.executive_summary.items():
        if isinstance(value, dict):
            continue
        story.append(Paragraph(f"{key.replace('_', ' ').title()}: {value}", styles["Normal"]))
    story.append(Spacer(1, 16))

    for section in report.sections:
        story.append(Paragraph(f"{section.title} ({section.status.value})", styles["Heading2"]))
        story.append(Paragraph(section.summary, styles["Normal"]))
        if section.details:
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(d, styles["Normal"])) for d in section.details[:50]],
                    bulletType="bullet",
                )
            )
        story.append(Spacer(1, 10))

    doc.build(story)
    return buffer.getvalue()


def render_docx(report: AssessmentReport) -> bytes:
    document = Document()
    document.add_heading(f"Migration Assessment Report — {report.project_name}", level=0)
    document.add_paragraph(f"Generated {report.generated_at.isoformat()}")

    document.add_heading("Executive Summary", level=1)
    for key, value in report.executive_summary.items():
        if isinstance(value, dict):
            continue
        document.add_paragraph(f"{key.replace('_', ' ').title()}: {value}")

    for section in report.sections:
        document.add_heading(f"{section.title} ({section.status.value})", level=2)
        document.add_paragraph(section.summary)
        for detail in section.details[:50]:
            paragraph = document.add_paragraph(detail, style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
